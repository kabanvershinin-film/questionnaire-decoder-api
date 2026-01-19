import json, re, os
from difflib import get_close_matches
from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from docx import Document
from io import BytesIO

app = FastAPI()

# CORS для GitHub Pages
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

def norm(s: str) -> str:
    s = s.strip()
    s = re.sub(r"\s+", " ", s)
    s = s.replace("\u00a0", " ").replace("\u200b", " ")
    s = re.sub(r"\([^)]*\)", "", s)  # убираем пояснения в скобках
    s = s.strip().lower()
    return s

# Загружаем базу расшифровок
if os.path.exists("decode_map.json"):
    with open("decode_map.json", "r", encoding="utf-8") as f:
        raw = json.load(f)
else:
    raw = {}

DECODE_MAP = {norm(k): v for k, v in raw.items()}
DECODE_KEYS = list(DECODE_MAP.keys())

PLUS_MARKS = {"+", "＋"}  # плюс обычный и полноширинный
TEXT_MARKS = ("☑", "☒", "✔", "X", "x")  # если где-то будут галочки в тексте

# насколько “похожим” должен быть пункт, чтобы считать совпадением
FUZZY_CUTOFF = 0.62

@app.get("/health")
def health():
    return {"status": "ok"}

@app.post("/decode")
async def decode(file: UploadFile = File(...)):
    data = await file.read()
    doc = Document(BytesIO(data))

    extracted = []

    # 1) Плюсы в таблицах (главное)
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            left = (row.cells[0].text or "").strip()
            right = (row.cells[1].text or "").strip()
            # плюс может быть с пробелами/переносами
            if (("+" in right) or ("＋" in right)) and left:
                extracted.append(left)

    # 2) На всякий случай: если где-то отметки стоят в обычном тексте
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        if "+" in t or "＋" in t or any(m in t for m in TEXT_MARKS):
            clean = t.replace("+", " ").replace("＋", " ")
            for m in TEXT_MARKS:
                clean = clean.replace(m, " ")
            clean = clean.strip(" -–—\t")
            if clean:
                extracted.append(clean)

    # Убираем дубли
    extracted = list(dict.fromkeys(extracted))

    matched, missed, report_parts = [], [], []
    missed_suggestions = {}  # что было “похоже” (для отладки)

    for symptom in extracted:
        key = norm(symptom)

        # 1) точное совпадение
        text = DECODE_MAP.get(key)

        # 2) если нет — пытаемся найти похожий ключ
        if not text and DECODE_KEYS:
            cand = get_close_matches(key, DECODE_KEYS, n=1, cutoff=FUZZY_CUTOFF)
            if cand:
                text = DECODE_MAP[cand[0]]
                missed_suggestions[symptom] = cand[0]  # покажем, что сопоставили

        if text:
            matched.append(symptom)
            report_parts.append(f"### {symptom}\n{text}")
        else:
            missed.append(symptom)

    report = "\n\n".join(report_parts) if report_parts else "Отмеченные пункты не найдены в базе расшифровок."

    return {
        "matched": matched,
        "missed": missed,
        "extracted": extracted,
        "missed_suggestions": missed_suggestions,
        "report_markdown": report
    }
